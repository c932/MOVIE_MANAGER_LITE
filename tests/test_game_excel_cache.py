"""
gamewall.excel_parser / gamewall.game_cache 单元测试
"""
import json
import os
import tempfile
from pathlib import Path

from openpyxl import Workbook

from gamewall.excel_parser import parse_excel, ExcelParseError
from gamewall.game_cache import GameCacheManager
from gamewall.game_models import Game


def _make_excel(path: Path):
    wb = Workbook()
    ws = wb.active
    ws['A1'] = '月下游戏资源合集'
    ws['A3'] = '免责声明...'
    ws['A8'] = '解压密码：YXGAME'
    ws.append([]); ws.append([]); ws.append([])
    # 表头行（第 8 行之后）
    ws.append(['游戏名（按CTRL+F搜索）', '下载链接（夸克）', '大小', '版本',
               '更新日期', '发布日期', 'IGN评分', 'Steam评分', '3A大作'])
    ws.append(['黎明行者之血/The Blood of Dawnwalker', 'https://pan.quark.cn/s/abc',
               '48.4G', 'v1.0.1 中文版', '2026-09-04', '', '', '', ''])
    ws.append(['艾尔登法环/ELDEN RING', 'https://pan.quark.cn/s/def',
               '59.5G', 'Build.24831693', '2026-09-03', '2022-02-25', '9.6', '92%', '是'])
    ws.append(['不心动才有鬼', 'https://pan.quark.cn/s/ghi', '506M', '中文版',
               '2026-09-01', '', '', '', ''])
    wb.save(path)


def test_parse_excel():
    import sys
    if sys.platform != 'win32':
        return  # 该用例依赖真实 Excel，仅在 Windows 跑
    real = r'E:\Downloads\Game\端游资源库9.4.xlsx'
    if not Path(real).exists():
        return
    games, meta = parse_excel(real)
    assert len(games) == 4694, f"实际解析 {len(games)}"
    assert meta['header_row'] == 12
    g0 = games[0]
    assert g0.title_cn == '黎明行者之血'
    assert g0.title_en == 'The Blood of Dawnwalker'
    assert g0.quark_link.startswith('https://pan.quark.cn/')
    assert g0.size_bytes > 48 * 1024 ** 3
    # 无斜杠名
    assert games[2].title_cn == '不心动才有鬼' and games[2].title_en == ''


def test_parse_excel_synthetic(tmp_path=None):
    with tempfile.TemporaryDirectory() as td:
        xlsx = Path(td) / 'games.xlsx'
        _make_excel(xlsx)
        games, meta = parse_excel(str(xlsx))
        assert meta['header_row'] == 12
        assert len(games) == 3
        g1 = games[1]
        assert g1.title_cn == '艾尔登法环'
        assert g1.title_en == 'ELDEN RING'
        assert g1.release_date == '2022-02-25'
        assert g1.ign_score == 9.6
        assert g1.steam_positive_pct == 92.0
        assert g1.aaa_excel_marked is True
        assert g1.is_aaa is True
        assert games[0].aaa_excel_marked is False
        assert games[0].is_aaa is False
        assert games[0].size_bytes == 48.4 * 1024 ** 3
        assert games[2].size_bytes == 506 * 1024 ** 2


def _cache(tmp: str) -> GameCacheManager:
    os.environ['APP_DIR_NAME_OVERRIDE'] = 'LocalGameWall'
    import utils.app_paths as ap
    ap.DATA_DIR = Path(tmp)
    return GameCacheManager()


def test_cache_merge_never_overwrite():
    with tempfile.TemporaryDirectory() as td:
        cache = _cache(td)
        g = Game(raw_name='艾尔登法环/ELDEN RING')
        # 联网补全：填空
        cache.merge_fields(g.norm_key, {'steam_appid': 1245620, 'description': 'xx',
                                        'steam_positive_pct': 90.0})
        # 再次联网：不覆盖已有值
        cache.merge_fields(g.norm_key, {'steam_appid': 999, 'steam_positive_pct': 10.0,
                                        'metacritic': 94.0})
        entry = cache.games_data[g.norm_key]
        assert entry['steam_appid'] == 1245620
        assert entry['steam_positive_pct'] == 90.0
        assert entry['metacritic'] == 94.0


def test_cache_excel_force_and_apply():
    with tempfile.TemporaryDirectory() as td:
        cache = _cache(td)
        g = Game(raw_name='艾尔登法环/ELDEN RING', ign_score=0.0)
        cache.merge_fields(g.norm_key, {'ign_score': 9.0, 'release_date': '2022-02-25'})
        # 用户在 Excel 填了 IGN → 强制覆盖缓存
        g2 = Game(raw_name='艾尔登法环/ELDEN RING', ign_score=9.6)
        cache.merge_fields(g2.norm_key, {'ign_score': 9.6}, force=True)
        assert cache.games_data[g.norm_key]['ign_score'] == 9.6
        # apply_to_game(force=False) 填充联网/本地字段的空值；评分可来自网页或用户手填，缓存值回填
        g3 = Game(raw_name='艾尔登法环/ELDEN RING')
        cache.apply_to_game(g3)
        assert g3.release_date == '2022-02-25'
        assert g3.ign_score == 9.6
        # Excel 字段不回填：is_aaa 保持 False（防止用户取消 3A 标记后被缓存复活）
        assert g3.is_aaa is False
        assert g3.quark_link == ''
        # 缓存优先降级路径：直接从缓存构建
        games = cache.load_games_from_cache()
        assert len(games) == 1
        assert games[0].release_date == '2022-02-25'


def test_cache_persistence():
    with tempfile.TemporaryDirectory() as td:
        cache = _cache(td)
        g = Game(raw_name='测试游戏/Test Game', quark_link='https://x.cn/s/1')
        cache.merge_fields(g.norm_key, {'steam_appid': 42, 'genres': ['动作', 'RPG']})
        cache.save()
        cache2 = GameCacheManager()
        assert len(cache2) == 1
        g2 = Game(raw_name='测试游戏/Test Game')
        cache2.apply_to_game(g2)
        assert g2.steam_appid == 42
        assert g2.genres == ['动作', 'RPG']


def test_scraper_neg_cache():
    with tempfile.TemporaryDirectory() as td:
        cache = _cache(td)
        g = Game(raw_name='某游戏/Some Game')
        assert cache.get_scraper_neg(g.norm_key, 'gamersky') is None
        cache.set_scraper_neg(g.norm_key, 'gamersky')
        assert cache.get_scraper_neg(g.norm_key, 'gamersky') is not None
        cache.save()
        cache2 = GameCacheManager()
        assert cache2.get_scraper_neg(g.norm_key, 'gamersky') is not None


def test_needs_enrichment():
    g = Game(raw_name='X/Y')
    assert GameCacheManager.needs_enrichment(g)
    g.steam_appid = 1
    g.description = 'd'
    g.release_date = '2020-01-01'
    g.steam_positive_pct = 90.0
    assert not GameCacheManager.needs_enrichment(g)


def test_enrichment_worker_tracks_all_steam_detail_fields():
    from gamewall.enrichment_worker import EnrichmentWorker

    game = Game(
        raw_name='测试游戏/Test Game',
        description='简介',
        release_date='2025-01-01',
        genres=['动作'],
        header_image='https://example.test/cover.jpg',
        metacritic=80.0,
        publishers=['Publisher'],
        platforms=['windows'],
        steam_categories=['Single-player'],
    )

    worker = EnrichmentWorker([game], enable_gamersky=False)

    assert worker._todo[0]['missing_details'] is False


def test_worker_runs_third_party_stage_without_steam_for_core_complete_game():
    from gamewall.enrichment_worker import EnrichmentWorker

    game = Game(
        raw_name='Steam 核心完整/Steam Complete',
        steam_appid=123,
        description='简介',
        release_date='2025-01-01',
        steam_positive_pct=90.0,
    )

    worker = EnrichmentWorker(
        [game],
        steam_candidates=set(),
        gamersky_candidates={game.norm_key},
        enable_ign=False,
    )
    steam_calls = []
    rating_calls = []

    class _Gamersky:
        def fetch_score(self, title):
            rating_calls.append(title)
            return 8.8, ''

    def record_steam_call(items):
        steam_calls.extend(items)
        return False

    worker._gamersky = _Gamersky()
    worker._run_steam_pipeline = record_steam_call
    worker.run()

    assert worker._todo[0]['steam_eligible'] is False
    assert steam_calls == []
    assert rating_calls == [game.raw_name]


def test_worker_records_steam_match_negative_immediately(monkeypatch):
    from gamewall import steam_client
    from gamewall.enrichment_worker import EnrichmentWorker

    game = Game(raw_name='确定不在Steam/Definitely Not On Steam')
    worker = EnrichmentWorker([game], enable_gamersky=False)
    negs = []
    worker.negative_recorded.connect(lambda key, source: negs.append((key, source)))
    monkeypatch.setattr(steam_client, 'match_appid', lambda queries: (0, '', True))

    worker._run_steam_pipeline(worker._todo)

    assert negs == [(game.norm_key, 'steam_match')]


def test_worker_skips_negative_when_steam_search_unreachable(monkeypatch):
    from gamewall import steam_client
    from gamewall.enrichment_worker import EnrichmentWorker

    game = Game(raw_name='断网候选/Offline Candidate')
    worker = EnrichmentWorker([game], enable_gamersky=False)
    negs = []
    worker.negative_recorded.connect(lambda key, source: negs.append((key, source)))
    monkeypatch.setattr(steam_client, 'match_appid', lambda queries: (0, '', False))

    worker._run_steam_pipeline(worker._todo)

    assert negs == []


def test_worker_records_third_party_negative_immediately():
    from gamewall.enrichment_worker import EnrichmentWorker

    game = Game(raw_name='第三方无分/No Third Party Score')
    worker = EnrichmentWorker(
        [game], steam_candidates=set(),
        gamersky_candidates={game.norm_key}, enable_ign=False)

    class _Empty:
        def fetch_score(self, title):
            return None, ''

    worker._gamersky = _Empty()
    negs = []
    worker.negative_recorded.connect(lambda key, source: negs.append((key, source)))
    worker.run()

    assert negs == [(game.norm_key, 'gamersky')]


def test_aaa_override_persists_across_snapshot_and_restart():
    from gamewall.aaa_classifier import classify_aaa

    with tempfile.TemporaryDirectory() as td:
        cache = _cache(td)
        game = Game(raw_name='测试游戏/Test Game', aaa_excel_marked=True)
        cache.snapshot_games([game])
        result = classify_aaa({
            'large_developer': {'confirmed': True, 'source': '官方开发团队介绍'},
        })
        cache.set_auto_aaa_classification(game.norm_key, result, {
            'large_developer': {'confirmed': True, 'source': '官方开发团队介绍'},
        })
        cache.set_manual_aaa_override(game.norm_key, False)
        cache.save()

        cache2 = GameCacheManager()
        updated_excel_game = Game(raw_name='测试游戏/Test Game', aaa_excel_marked=False)
        cache2.apply_to_game(updated_excel_game)
        assert updated_excel_game.aaa_manual_override is False
        assert updated_excel_game.aaa_score == 15
        assert updated_excel_game.is_aaa is False

        cache2.snapshot_games([updated_excel_game])
        cache2.save()
        cache3 = GameCacheManager()
        restored = Game(raw_name='测试游戏/Test Game', aaa_excel_marked=False)
        cache3.apply_to_game(restored)
        assert restored.aaa_manual_override is False
        assert restored.is_aaa is False

        cache3.clear_manual_aaa_override(restored.norm_key)
        assert cache3.get_manual_aaa_override(restored.norm_key) is None
        cache3.apply_to_game(restored)
        restored.aaa_manual_override = None
        assert restored.is_aaa is False


def test_v1_cache_does_not_promote_legacy_aaa_flag():
    with tempfile.TemporaryDirectory() as td:
        _cache(td)
        cache_path = Path(td) / 'game_cache.json'
        cache_path.write_text(json.dumps({
            'version': 1,
            'games': {'legacy': {'is_aaa': True, 'steam_appid': 42}},
            'scraper_neg': {},
        }), encoding='utf-8')
        cache = GameCacheManager()
        assert cache.games_data['legacy']['steam_appid'] == 42
        assert 'is_aaa' not in cache.games_data['legacy']
        legacy = Game(raw_name='legacy')
        cache.apply_to_game(legacy)
        assert legacy.aaa_manual_override is None
        assert legacy.is_aaa is False


if __name__ == '__main__':
    test_parse_excel()
    test_parse_excel_synthetic()
    test_cache_merge_never_overwrite()
    test_cache_excel_force_and_apply()
    test_cache_persistence()
    test_scraper_neg_cache()
    test_needs_enrichment()
    test_aaa_override_persists_across_snapshot_and_restart()
    test_v1_cache_does_not_promote_legacy_aaa_flag()
    print('all excel/cache tests passed')
